"""
Shared Excel report generator used by Web, Android, and iOS engines.
Place this file at: core/excel_reporter.py
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelReporter:

    HEADER_FILL  = PatternFill("solid", start_color="1F2937", end_color="1F2937")
    HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11, name="Arial")
    PASS_FILL    = PatternFill("solid", start_color="D1FAE5", end_color="D1FAE5")
    FAIL_FILL    = PatternFill("solid", start_color="FEE2E2", end_color="FEE2E2")
    TITLE_FONT   = Font(bold=True, size=16, name="Arial", color="1F2937")
    SUB_FONT     = Font(size=10, name="Arial", color="6B7280")
    LABEL_FONT   = Font(bold=True, size=10, name="Arial")
    SECTION_FONT = Font(bold=True, size=11, name="Arial", color="1F2937")
    THIN_BORDER  = Border(*[Side(style="thin", color="E5E7EB")] * 4)

    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
        self.timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # ==========================
    # MAIN EXPORT
    # ==========================
    def export(self, platform: str, target: str, results: list,
               discovery_summary: dict = None) -> str:

        wb = Workbook()

        self._build_summary_sheet(wb, platform, target, results, discovery_summary)
        self._build_discovery_sheet(wb, platform, discovery_summary)
        self._build_testcases_sheet(wb, results)
        self._build_bugs_sheet(wb, results)

        path = os.path.join(
            self.output_dir,
            f"{platform.lower()}_test_report_{self.timestamp}.xlsx"
        )
        wb.save(path)
        return path

    # ==========================
    # SHEET 1 — SUMMARY
    # ==========================
    def _build_summary_sheet(self, wb, platform, target, results, discovery_summary):
        ws = wb.active
        ws.title = "Summary"

        passed = [r for r in results if r["status"] == "PASS"]
        failed = [r for r in results if r["status"] == "FAIL"]
        total  = len(results)
        rate   = round(len(passed) / max(total, 1) * 100, 1)

        ws["B2"] = "AUTOMATION TEST REPORT"
        ws["B2"].font = self.TITLE_FONT
        ws["B3"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["B3"].font = self.SUB_FONT

        rows = [
            ("Platform",   platform),
            ("Target",     target),
            ("Total Tests", total),
            ("Passed",      len(passed)),
            ("Failed",      len(failed)),
            ("Pass Rate",   f"{rate}%"),
        ]
        r = 5
        for label, value in rows:
            ws.cell(row=r, column=2, value=label).font = self.LABEL_FONT
            cell = ws.cell(row=r, column=3, value=value)
            if label == "Passed":
                cell.font = Font(bold=True, color="059669")
            elif label == "Failed":
                cell.font = Font(bold=True, color="DC2626")
            r += 1

        # Category breakdown
        categories = {}
        for res in results:
            cat = res.get("category", "general")
            categories.setdefault(cat, {"pass": 0, "fail": 0})
            categories[cat]["pass" if res["status"] == "PASS" else "fail"] += 1

        r += 1
        ws.cell(row=r, column=2, value="CATEGORY BREAKDOWN").font = self.LABEL_FONT
        r += 1
        ws.cell(row=r, column=2, value="Category").font = self.HEADER_FONT
        ws.cell(row=r, column=2).fill = self.HEADER_FILL
        ws.cell(row=r, column=3, value="Pass").font = self.HEADER_FONT
        ws.cell(row=r, column=3).fill = self.HEADER_FILL
        ws.cell(row=r, column=4, value="Fail").font = self.HEADER_FONT
        ws.cell(row=r, column=4).fill = self.HEADER_FILL
        r += 1
        for cat, counts in categories.items():
            ws.cell(row=r, column=2, value=cat.replace("_", " ").title())
            ws.cell(row=r, column=3, value=counts["pass"])
            ws.cell(row=r, column=4, value=counts["fail"])
            r += 1

        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 12

    # ==========================
    # SHEET 2 — DISCOVERY DETAILS  (NEW)
    # ==========================
    def _build_discovery_sheet(self, wb, platform: str, discovery_summary: dict):
        ws = wb.create_sheet("Discovery Details")

        if not discovery_summary:
            ws["B2"] = "No discovery data captured for this run."
            ws["B2"].font = Font(italic=True, color="6B7280", size=11)
            ws.column_dimensions["B"].width = 40
            return

        row = 2

        # ── Page / App Info block ──
        ws.cell(row=row, column=2, value="PAGE / APP INFO").font = self.SECTION_FONT
        row += 1

        info_map = self._extract_info_fields(platform, discovery_summary)
        for label, value in info_map:
            ws.cell(row=row, column=2, value=label).font = self.LABEL_FONT
            ws.cell(row=row, column=3, value=value)
            row += 1

        row += 1

        # ── Inputs table ──
        inputs = discovery_summary.get("input_details", [])
        row = self._write_element_table(
            ws, row, f"INPUT FIELDS ({len(inputs)})",
            inputs, self._input_columns(platform)
        )

        row += 1

        # ── Buttons table ──
        buttons = discovery_summary.get("button_details", [])
        row = self._write_element_table(
            ws, row, f"BUTTONS ({len(buttons)})",
            buttons, self._button_columns(platform)
        )

        row += 1

        # ── Links / Clickables / Tappables table ──
        if platform == "WEB":
            links = discovery_summary.get("link_details", [])
            row = self._write_element_table(
                ws, row, f"LINKS ({len(links)})",
                links, [("Text", "text"), ("Href", "href")]
            )
        elif platform == "ANDROID":
            clickables = discovery_summary.get("clickable_details", [])
            row = self._write_element_table(
                ws, row, f"CLICKABLE ELEMENTS ({len(clickables)})",
                clickables, [("Class", "class"), ("Text", "text"), ("Resource ID", "resource_id")]
            )
        elif platform == "IOS":
            tappables = discovery_summary.get("tappable_details", [])
            row = self._write_element_table(
                ws, row, f"TAPPABLE ELEMENTS ({len(tappables)})",
                tappables, [("Type", "type"), ("Label", "label")]
            )

        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 28
        ws.column_dimensions["E"].width = 28
        ws.freeze_panes = "A2"

    def _extract_info_fields(self, platform: str, d: dict) -> list:
        if platform == "WEB":
            return [
                ("Page Title",   d.get("title", "")),
                ("URL",          d.get("url", "")),
                ("Total Inputs", d.get("inputs", 0)),
                ("Total Buttons", d.get("buttons", 0)),
                ("Total Links",  d.get("links", 0)),
            ]
        if platform == "ANDROID":
            return [
                ("Package",      d.get("package", "")),
                ("Activity",     d.get("activity", "")),
                ("Total Inputs", d.get("inputs", 0)),
                ("Total Buttons", d.get("buttons", 0)),
                ("Total Clickables", d.get("clickables", 0)),
            ]
        if platform == "IOS":
            return [
                ("Bundle ID",        d.get("bundle_id", "")),
                ("Text Fields",      d.get("text_fields", 0)),
                ("Secure Fields",    d.get("secure_fields", 0)),
                ("Total Buttons",    d.get("buttons", 0)),
                ("Cells",            d.get("cells", 0)),
                ("Static Texts",     d.get("static_texts", 0)),
            ]
        return []

    def _input_columns(self, platform: str) -> list:
        if platform == "WEB":
            return [("Type", "type"), ("Name", "name"), ("Placeholder", "placeholder"), ("ID", "id")]
        if platform == "ANDROID":
            return [("Text", "text"), ("Resource ID", "resource_id"), ("Hint", "hint"), ("Enabled", "enabled")]
        if platform == "IOS":
            return [("Type", "type"), ("Label", "label"), ("Placeholder", "placeholder"), ("Enabled", "enabled")]
        return []

    def _button_columns(self, platform: str) -> list:
        if platform == "WEB":
            return [("Text", "text"), ("ID", "id"), ("Source", "source"), ("Enabled", "enabled")]
        if platform == "ANDROID":
            return [("Text", "text"), ("Resource ID", "resource_id"), ("Enabled", "enabled"), ("Clickable", "clickable")]
        if platform == "IOS":
            return [("Label", "label"), ("Enabled", "enabled"), ("Visible", "visible")]
        return []

    def _write_element_table(self, ws, row: int, title: str, items: list, columns: list) -> int:
        """Writes a titled, headered table of element dicts. Returns next free row."""

        ws.cell(row=row, column=2, value=title).font = self.SECTION_FONT
        row += 1

        if not items:
            ws.cell(row=row, column=2, value="None found").font = Font(italic=True, color="9CA3AF")
            return row + 2

        header_row = row
        for i, (label, _) in enumerate(columns):
            c = ws.cell(row=header_row, column=2 + i, value=label)
            c.font = self.HEADER_FONT
            c.fill = self.HEADER_FILL
            c.alignment = Alignment(horizontal="center")
        row += 1

        for item in items:
            for i, (_, key) in enumerate(columns):
                val = item.get(key, "")
                c = ws.cell(row=row, column=2 + i, value=str(val) if val != "" else "")
                c.border = self.THIN_BORDER
            row += 1

        return row + 1  # blank line after table

    # ==========================
    # SHEET 3 — ALL TEST CASES
    # ==========================
    def _build_testcases_sheet(self, wb, results: list):
        ws = wb.create_sheet("Test Cases")

        headers = ["ID", "Test Name", "Category", "Type", "Priority",
                   "Status", "Attempts", "Expected Result", "Error", "Timestamp"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = self.HEADER_FONT
            c.fill = self.HEADER_FILL
            c.alignment = Alignment(horizontal="center")

        for res in results:
            row = [
                res.get("id", ""),
                res.get("name", ""),
                res.get("category", "").replace("_", " ").title(),
                res.get("type", ""),
                res.get("priority", ""),
                res.get("status", ""),
                res.get("attempts", ""),
                res.get("expected", ""),
                res.get("error", "") or "",
                res.get("timestamp", ""),
            ]
            ws.append(row)
            r = ws.max_row
            fill = self.PASS_FILL if res.get("status") == "PASS" else self.FAIL_FILL
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).fill = fill
                ws.cell(row=r, column=col).border = self.THIN_BORDER

        widths = [8, 32, 16, 10, 10, 9, 9, 32, 40, 20]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # ==========================
    # SHEET 4 — BUGS ONLY
    # ==========================
    def _build_bugs_sheet(self, wb, results: list):
        bugs = [r for r in results if r["status"] == "FAIL"]
        ws = wb.create_sheet("Bugs Found")

        if not bugs:
            ws["B2"] = "No bugs found — all tests passed."
            ws["B2"].font = Font(bold=True, color="059669", size=12)
            return

        headers = ["Bug ID", "Test Name", "Category", "Priority",
                   "Attempts Before Confirmed", "Error Details", "Expected Result", "Found At"]
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = self.HEADER_FONT
            c.fill = self.HEADER_FILL
            c.alignment = Alignment(horizontal="center")

        for bug in bugs:
            ws.append([
                bug.get("id", ""),
                bug.get("name", ""),
                bug.get("category", "").replace("_", " ").title(),
                bug.get("priority", "").upper(),
                bug.get("attempts", ""),
                bug.get("error", ""),
                bug.get("expected", ""),
                bug.get("timestamp", ""),
            ])
            r = ws.max_row
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).fill = self.FAIL_FILL
                ws.cell(row=r, column=col).border = self.THIN_BORDER
                ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")

        widths = [9, 32, 16, 10, 14, 45, 32, 20]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"